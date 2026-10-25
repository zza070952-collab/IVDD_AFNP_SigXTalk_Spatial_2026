import collections,csv,hashlib,json,math,re
from pathlib import Path
import openpyxl
import argparse
ap=argparse.ArgumentParser();ap.add_argument('--project',type=Path,required=True);ap.add_argument('--revision',type=Path,required=True);ap.add_argument('--payload',type=Path,required=True);ap.add_argument('--original-workbook',type=Path,required=True);a=ap.parse_args()
R=a.project.resolve();C=a.revision.resolve();p=json.loads(a.payload.read_text(encoding='utf-8'))
def same(a,b):
 if a is None or b is None:return a is None and b is None
 if isinstance(a,bool) or isinstance(b,bool):return isinstance(a,bool) and isinstance(b,bool) and a==b
 if isinstance(a,(float,int)) and isinstance(b,(float,int)):return math.isclose(a,b,rel_tol=2e-13,abs_tol=2e-13)
 return a==b
reports=[]
if not p.get('books'):raise ValueError('No expected workbooks in payload')
key_fields={
 'S1_Cohorts_QC':['cohort_GEO','sample_GSM','study_role','donor_or_scope','tissue_or_identity','measurement'],
 'S2_Fixed_NP_Programs':['entry_type','cohort_or_scope','donor','program','gene_or_context','frozen_rank','measurement'],
 'S3_Official_SigXTalk':['record_type','record_key'],
 'S4_Spatial_Regions':['record_type','capture','disc','entity'],
 'S5_Cross_Evidence':['record_type','gene','source_id'],
 'S6_Oligos_and_Design':['Record_type','Record_ID'],
 'Raw_qPCR_Ct':['Source_sheet','Source_excel_row'],
 'Raw_CCK8_OD450':['Source_sheet','Source_excel_row'],
 'Batch_Results_Statistics':['Record_type','Batch_ID','Group_ID','siRNA','Gene','Contrast','Measure'],
}
for b in p['books']:
 f=C/'INTERNAL_REVIEW'/b['filename']
 if not f.is_file():raise FileNotFoundError('Expected workbook was not exported: '+str(f))
 wb=openpyxl.load_workbook(f,data_only=False,read_only=False)
 assert wb.sheetnames==[s['name'] for s in b['sheets']],wb.sheetnames
 report={'file':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'sheets':[]}
 for s in b['sheets']:
  ws=wb[s['name']];n=len(s['rows']);m=len(s['headers'])
  assert [ws.cell(5,j+1).value for j in range(m)]==s['headers']
  mismatch=[]
  for i,row in enumerate(s['rows'],6):
   for j,want in enumerate(row,1):
    got=ws.cell(i,j).value
    if not same(want,got):mismatch.append([i,j,want,got])
  assert not mismatch,(s['name'],mismatch[:6])
  kk=key_fields[s['name']];keys=[tuple(row[s['headers'].index(k)] for k in kk) for row in s['rows']]
  assert len(set(keys))==len(keys),(s['name'],'duplicate record keys')
  formulas=[c.coordinate for row in ws for c in row if c.data_type=='f'];assert not formulas
  assert not list(ws.merged_cells.ranges)
  assert ws.sheet_state=='visible'
  assert not any(x.hidden for x in ws.row_dimensions.values())
  assert not any(x.hidden for x in ws.column_dimensions.values())
  assert str(ws.freeze_panes)=='B6'
  assert len(ws.tables)==1
  assert next(iter(ws.tables.values())).ref==f'A5:{openpyxl.utils.get_column_letter(m)}{n+5}'
  text='\n'.join(str(c.value) for row in ws for c in row if c.value is not None)
  private=re.findall(r'[A-Za-z]:[/\\]Users[/\\]|/(?:home|Users)/[^/\s]+/|ghp_[A-Za-z0-9]{20,}|github_pat_[A-Za-z0-9_]{20,}',text);assert not private,private
  report['sheets'].append({'sheet':s['name'],'data_rows':n,'columns':m,'all_payload_cells_match':True,'unique_key_fields':kk,'duplicate_keys':0,'formulas':0,'merged_cells':0,'hidden_rows_or_columns':0,'freeze_panes':'B6','filter':True,'private_path_hits':private})
 # Raw data must equal the original supplied workbook, not just the derived cache.
 if f.name.startswith('Additional_file_3'):
  original=openpyxl.load_workbook(a.original_workbook,data_only=True,read_only=True)
  originals={}
  for ws in original:
   vals=list(ws.values);headers=list(vals[0]);originals[ws.title]=(headers,vals[1:])
  for sname,original_names in [('Raw_qPCR_Ct',['siRNA_qPCR_Raw','Program_qPCR_Raw']),('Raw_CCK8_OD450',['CCK8_Raw'])]:
   ws=wb[sname];headers=[c.value for c in ws[5]];index={x:i+1 for i,x in enumerate(headers)};seen=set();counts=collections.Counter()
   for row in range(6,6+next(len(s['rows']) for s in b['sheets'] if s['name']==sname)):
    source=ws.cell(row,index['Source_sheet']).value;source_row=ws.cell(row,index['Source_excel_row']).value
    assert source in original_names
    key=(source,source_row);assert key not in seen;seen.add(key);counts[source]+=1
    oh,ov=originals[source]
    for j,h in enumerate(oh):assert same(ov[source_row-2][j],ws.cell(row,index[h]).value),(sname,row,h)
   assert dict(counts)=={name:len(originals[name][1]) for name in original_names}
  qp=wb['Raw_qPCR_Ct'];qh=[c.value for c in qp[5]];recordtypes=collections.Counter(qp.cell(i,qh.index('Sample_type')+1).value for i in range(6,638));assert recordtypes=={'Sample':576,'NTC':24,'NoRT':32},recordtypes
  report['raw_original_comparison']={'qPCR_wells':632,'qPCR_sample':576,'qPCR_NTC':24,'qPCR_NoRT':32,'CCK8_wells':288,'all_original_cells_match':True,'unique_source_sheet_row_keys':True}
 reports.append(report)
(C/'audit/WORKBOOK_CONTENT_VALIDATION.json').write_text(json.dumps({'status':'PASS','reports':reports},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'validated_books':len(reports),'sheets':sum(len(r['sheets']) for r in reports),'status':'PASS'},indent=2))
