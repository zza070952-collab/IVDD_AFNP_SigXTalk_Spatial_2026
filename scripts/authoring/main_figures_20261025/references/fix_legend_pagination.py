#!/usr/bin/env python3
"""Apply only keepNext/keepLines to private DOCX legend groups.

The selected input is read-only. Export and inspect the resulting pages before
claiming visual approval; an exact XML difference check is not visual review.
"""
import argparse
import hashlib
from io import BytesIO
import json
from pathlib import Path
import zipfile

from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from output_guard import require_private_output


def sha(data):
    return hashlib.sha256(data).hexdigest()


def payloads(data):
    with zipfile.ZipFile(BytesIO(data)) as package:
        return {name: package.read(name) for name in package.namelist()}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True, help='Fresh private output DOCX')
    parser.add_argument('--expected-sha256', required=True)
    parser.add_argument('--audit', type=Path, required=True, help='Private JSON destination')
    parser.add_argument('--legend-heading', default='Figure legends')
    parser.add_argument('--end-heading', default='References')
    parser.add_argument('--expected-groups', type=int, default=6)
    args = parser.parse_args()
    target = require_private_output(args.output)
    audit_path = require_private_output(args.audit)
    if target == args.source.resolve() or audit_path == args.source.resolve() or audit_path == target:
        raise ValueError('Source, output and audit must be separate files')
    if target.exists() or audit_path.exists():
        raise FileExistsError('Choose fresh private output and audit files')
    before = args.source.read_bytes()
    if sha(before) != args.expected_sha256:
        raise ValueError('Unexpected selected input hash')
    doc = Document(BytesIO(before))
    paragraphs = doc.paragraphs
    heading = next(i for i, p in enumerate(paragraphs) if p.text == args.legend_heading)
    end = next(i for i, p in enumerate(paragraphs) if i > heading and p.text == args.end_heading)
    starts = [i for i in range(heading + 1, end) if paragraphs[i].text.startswith('Figure ')]
    if len(starts) != args.expected_groups:
        raise ValueError('Unexpected number of legend groups')
    all_text = [p.text for p in paragraphs]
    audit = []
    paragraphs[heading].paragraph_format.keep_with_next = True
    audit.append({'paragraph': heading, 'role': 'legend heading', 'keep_with_next': True})
    for k, start in enumerate(starts):
        stop = starts[k + 1] if k + 1 < len(starts) else end
        for i in range(start, stop):
            if not paragraphs[i].text:
                raise ValueError('Unexpected blank inside legend group')
            last = i == stop - 1
            paragraphs[i].paragraph_format.keep_together = True
            paragraphs[i].paragraph_format.keep_with_next = not last
            audit.append({'paragraph': i, 'legend': k + 1, 'keep_together': True, 'keep_with_next': not last})
    buffer = BytesIO()
    doc.save(buffer)
    after = buffer.getvalue()
    if [p.text for p in Document(BytesIO(after)).paragraphs] != all_text:
        raise ValueError('Layout patch changed paragraph text')
    original, updated = payloads(before), payloads(after)
    if set(original) != set(updated):
        raise ValueError('Package membership changed')
    changed = [name for name in original if original[name] != updated[name]]
    if set(changed) - {'word/document.xml'}:
        raise ValueError('A non-layout package part changed')
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    lhs, rhs = etree.fromstring(original['word/document.xml']), etree.fromstring(updated['word/document.xml'])
    for root in (lhs, rhs):
        ps = root.xpath('./w:body/w:p', namespaces=ns)
        for i in range(heading, end):
            properties = ps[i].find(qn('w:pPr'))
            if properties is not None:
                for tag in ('w:keepNext', 'w:keepLines'):
                    for element in list(properties.findall(qn(tag))):
                        properties.remove(element)
                if not len(properties):
                    ps[i].remove(properties)
    if etree.tostring(lhs) != etree.tostring(rhs):
        raise ValueError('XML changed outside authorized pagination tags')
    if args.source.read_bytes() != before:
        raise ValueError('Read-only source changed')
    target.parent.mkdir(parents=True, exist_ok=True)
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(after)
    receipt = {'status': 'PASS_LAYOUT_ONLY; VISUAL_QA_PENDING', 'input_sha256': sha(before),
               'output_sha256': sha(after), 'changed_package_parts': changed,
               'authorized_xml_tags': ['w:pPr/w:keepNext', 'w:pPr/w:keepLines'],
               'modified_paragraphs': audit, 'all_paragraph_text': 'EXACTLY_UNCHANGED',
               'all_media_and_other_package_parts': 'EXACTLY_UNCHANGED',
               'all_other_document_xml': 'EXACTLY_UNCHANGED_AFTER_REMOVING_ONLY_AUTHORIZED_TAGS',
               'source_input_unchanged': True, 'visual_QA': 'PENDING_NEW_EXPORT_AND_PAGE_REVIEW'}
    audit_path.write_text(json.dumps(receipt, indent=2), encoding='utf-8')
    print(json.dumps({'output_sha256': sha(after), 'patched_paragraphs': len(audit), 'visual_QA': 'PENDING'}))


if __name__ == '__main__':
    main()
