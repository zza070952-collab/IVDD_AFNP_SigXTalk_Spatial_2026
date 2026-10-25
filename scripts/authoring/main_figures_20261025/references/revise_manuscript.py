#!/usr/bin/env python3
"""Plan-driven private DOCX edit, preserving unedited text and source files.

The public code contains no study prose, result values, legends or source
document hash. Supply those inputs in a private plan and legend file. Outputs
contain manuscript text and must remain outside the publication checkout.
"""
import argparse
from copy import deepcopy
import csv
import hashlib
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph

from figure_references import extract, rewrite, scientific_numbers
from output_guard import require_private_output


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def set_text(paragraph, value):
    first = next((r for r in paragraph.runs if r.text), None)
    style = deepcopy(first._r.rPr) if first is not None and first._r.rPr is not None else None
    for child in list(paragraph._p):
        if child.tag != qn('w:pPr'):
            paragraph._p.remove(child)
    run = paragraph.add_run(value)
    if style is not None:
        run._r.insert(0, style)


def replace_spans(paragraph, spans):
    """Apply nonoverlapping spans without rewriting unaffected runs."""
    expected = paragraph.text
    nodes = list(paragraph._p.iter(qn('w:t')))
    for start, end, replacement in sorted(spans, reverse=True):
        offsets, cursor = [], 0
        for node in nodes:
            value = node.text or ''
            offsets.append((cursor, cursor + len(value), node, value))
            cursor += len(value)
        if cursor != len(expected):
            raise ValueError('Unsupported non-text paragraph elements')
        affected = [x for x in offsets if x[0] < end and x[1] > start]
        if not affected:
            raise ValueError('Missing replacement span')
        for j, (a, b, node, value) in enumerate(affected):
            before = value[:max(0, start - a)] if j == 0 else ''
            after = value[max(0, end - a):] if j == len(affected) - 1 else ''
            node.text = before + (replacement if j == 0 else '') + after
            node.set(qn('xml:space'), 'preserve')
        expected = expected[:start] + replacement + expected[end:]
    if paragraph.text != expected:
        raise ValueError('OOXML text edit did not round-trip')


def drop_panel(text, panel):
    matches = list(re.finditer(r'\(([A-G])\)\s', text))
    if [m.group(1) for m in matches].count(panel) != 1:
        raise ValueError('Expected exactly one requested panel marker')
    hit = next(i for i, m in enumerate(matches) if m.group(1) == panel)
    begin = matches[hit].start()
    end = matches[hit + 1].start() if hit + 1 < len(matches) else len(text)
    return text[:begin] + text[end:], text[begin:end]


def supplementary_refs(text):
    return re.findall(r'(?:Supplementary\s+)?(?:Figures?|Tables?)\s+S\d+(?:[A-Z]|[–-]S?\d+)?', text)


def add_before(paragraph, value):
    node = OxmlElement('w:p')
    paragraph._p.addprevious(node)
    result = Paragraph(node, paragraph._parent)
    if paragraph._p.pPr is not None:
        result._p.append(deepcopy(paragraph._p.pPr))
    result.add_run(value)
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        result.runs[0]._r.insert(0, deepcopy(paragraph.runs[0]._r.rPr))
    return result


def tsv(path, rows, columns):
    with Path(path).open('w', newline='', encoding='utf-8-sig') as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, delimiter='\t', extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)


def validated_plan(path):
    plan = json.loads(Path(path).read_text(encoding='utf-8'))
    required = ['source_sha256', 'expected_paragraph_count', 'edit_stop', 'preserve_ranges',
                'structure_checks', 'panel_edits', 'reference_overrides', 'text_replacements',
                'insertions', 'old_images_start', 'expected_old_images',
                'reject_old_image_indices', 'expected_sections', 'required_final_phrases',
                'forbidden_final_phrases', 'max_panel', 'legend_section']
    if any(key not in plan for key in required):
        raise ValueError('Private plan is incomplete; see plan.schema.json')
    if not re.fullmatch(r'[0-9a-f]{64}', plan['source_sha256']):
        raise ValueError('Private plan requires the selected source SHA-256')
    if set(plan['max_panel']) != {str(n) for n in range(1, 7)}:
        raise ValueError('Exactly six final main figures are required')
    if sum(item['source'] == 'figure1_legend' for item in plan['insertions']) != 1:
        raise ValueError('Exactly one private Figure 1 legend insertion is required')
    for item in plan['insertions']:
        if item['source'] not in ('figure1_legend', 'text') or not 0 <= item['before_index'] < plan['edit_stop']:
            raise ValueError('Invalid private-plan insertion')
        if item['source'] == 'text' and not isinstance(item.get('text'), str):
            raise ValueError('Text insertion requires an external text value')
    name = plan.get('output_name', 'manuscript_renumbered.docx')
    if Path(name).name != name or not name.lower().endswith('.docx'):
        raise ValueError('output_name must be a DOCX basename')
    plan['output_name'] = name
    return plan


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True, help='Private destination outside the publication checkout')
    parser.add_argument('--plan', type=Path, required=True, help='Private JSON with all source-specific text, assertions and indices')
    parser.add_argument('--figure-manifest', type=Path, required=True, help='Private JSON {"1":"image.png",...,"6":"image.png"}')
    parser.add_argument('--figure1-legend', type=Path, required=True, help='Private UTF-8 legend text')
    args = parser.parse_args()
    plan = validated_plan(args.plan)
    source = args.source.resolve()
    output = require_private_output(args.output_dir)
    if source.parent == output or source.parent in output.parents:
        raise ValueError('Output must be independent of the read-only source directory')
    if sha(source) != plan['source_sha256']:
        raise ValueError('Selected source SHA mismatch')
    figures = json.loads(args.figure_manifest.read_text(encoding='utf-8'))
    if set(figures) != {str(x) for x in range(1, 7)}:
        raise ValueError('Exactly six final figure files are required')
    figures = {int(n): Path(path).resolve() for n, path in figures.items()}
    if any(not p.is_file() or p.suffix.lower() != '.png' for p in figures.values()):
        raise ValueError('Each final embedded figure must be a real PNG')
    legend = args.figure1_legend.read_text(encoding='utf-8').strip()
    expected_panels = list('ABCDE')
    if not legend.startswith('Figure 1') or re.findall(r'\(([A-E])\)', legend) != expected_panels:
        raise ValueError('Figure 1 legend must have exactly A–E in artwork order')
    doc = Document(source)
    paragraphs = list(doc.paragraphs)
    if len(paragraphs) != plan['expected_paragraph_count']:
        raise ValueError('Unexpected input paragraph count')
    for check in plan['structure_checks']:
        value = paragraphs[check['index']].text
        if ('equals' in check and value != check['equals']) or ('starts_with' in check and not value.startswith(check['starts_with'])):
            raise ValueError('Input structure assertion failed at paragraph ' + str(check['index']))
    before = {i: p.text for i, p in enumerate(paragraphs)}
    before_xml = {i: p._p.xml for i, p in enumerate(paragraphs)}
    preserve = {i for start, stop in plan['preserve_ranges'] for i in range(start, stop)}
    if not (0 < plan['edit_stop'] <= plan['old_images_start'] <= len(paragraphs)):
        raise ValueError('Invalid private plan edit/image ranges')
    panels = {x['index']: x for x in plan['panel_edits']}
    changes, references, deleted, replacements, guards = [], [], [], [], []
    for i, paragraph in enumerate(paragraphs[:plan['edit_stop']]):
        old = paragraph.text
        if i in preserve:
            continue
        guard_source = old
        if i in panels:
            spec = panels[i]
            guard_source, removed = drop_panel(old, spec['remove_panel'])
            deleted.append({'old_paragraph': i, 'old_panel': spec['old_panel'], 'removed_text': removed, 'reason': spec['reason']})
            new, rows = rewrite(guard_source)
            relabel = spec.get('panel_relabel', {})
            new = re.sub(r'\(([A-G])\)(?=\s)', lambda m: '(' + relabel.get(m.group(1), m.group(1)) + ')', new)
            set_text(paragraph, new)
        else:
            spans, rows = [], []
            for match in extract(old):
                key = str(i) + '|' + match['text']
                override = plan['reference_overrides'].get(key)
                if override is not None:
                    newref = override['new_ref']
                    row = {'old_ref': match['text'], 'new_ref': newref, 'reason': override['reason']}
                else:
                    newref, mapped = rewrite(match['text'])
                    row = mapped[0]
                spans.append((match['start'], match['end'], newref))
                rows.append(row)
            replace_spans(paragraph, spans)
        for spec in plan['text_replacements'].get(str(i), []):
            old_value, new_value = spec['old'], spec['new']
            if not old_value or paragraph.text.count(old_value) != 1:
                raise ValueError('Expected exactly one configured replacement at paragraph ' + str(i))
            begin = paragraph.text.index(old_value)
            replace_spans(paragraph, [(begin, begin + len(old_value), new_value)])
            replacements.append({'old_paragraph': i, **spec})
        new = paragraph.text
        if scientific_numbers(guard_source) != scientific_numbers(new):
            raise ValueError('Scientific numeric drift at paragraph ' + str(i))
        if supplementary_refs(old) != supplementary_refs(new):
            raise ValueError('Supplementary reference drift at paragraph ' + str(i))
        guards.append({'old_paragraph': i, 'status': 'PASS', 'before_tokens': scientific_numbers(guard_source),
                       'after_tokens': scientific_numbers(new), 'deleted_panel_exception': i in panels})
        for row in rows:
            references.append({'old_paragraph': i, **row})
        if old != new:
            changes.append({'old_paragraph': i, 'operation': 'configured_figure_edit', 'old_text': old, 'new_text': new, 'numeric_guard': 'PASS'})
    inserted = []
    for spec in plan['insertions']:
        index = spec['before_index']
        value = legend if spec['source'] == 'figure1_legend' else spec['text']
        add_before(paragraphs[index], value)
        inserted.append({'before_index': index, 'text': value})
        changes.append({'old_paragraph': 'INSERT_BEFORE_' + str(index), 'operation': spec['source'],
                        'old_text': '', 'new_text': value, 'numeric_guard': 'EXPLICIT_INSERTION'})
        for reference in extract(value):
            references.append({'old_paragraph': 'INSERT_BEFORE_' + str(index), 'old_ref': 'NEW', 'new_ref': reference['text'], 'reason': 'Explicit private-plan insertion'})
    old_rel_ids = []
    for paragraph in paragraphs[plan['old_images_start']:]:
        old_rel_ids.extend(paragraph._p.xpath('.//a:blip/@r:embed'))
        paragraph._p.getparent().remove(paragraph._p)
    old_hashes = []
    for rid in dict.fromkeys(old_rel_ids):
        old_hashes.append(hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest())
        doc.part.drop_rel(rid)
    if len(old_hashes) != plan['expected_old_images']:
        raise ValueError('Unexpected original embedded-image count')
    section = doc.sections[-1]
    max_width = min(int(Mm(180)), section.page_width - section.left_margin - section.right_margin)
    max_height = section.page_height - section.top_margin - section.bottom_margin - int(Mm(15))
    embedded = []
    for n in range(1, 7):
        heading = doc.add_paragraph(plan.get('figure_headings', {}).get(str(n), 'Figure ' + str(n)))
        heading.paragraph_format.page_break_before = n > 1
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_after = Pt(6)
        heading.runs[0].bold = True
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        from PIL import Image
        with Image.open(figures[n]) as image:
            width, height = image.size
        scale = min(max_width / width, max_height / height)
        inline = paragraph.add_run().add_picture(str(figures[n]), width=int(width * scale), height=int(height * scale))
        inline._inline.docPr.set('descr', plan.get('figure_alt_text', {}).get(str(n), 'Figure ' + str(n)))
        embedded.append({'figure': n, 'sha256': sha(figures[n]), 'width_mm': width * scale / 36000,
                         'height_mm': height * scale / 36000})
        references.append({'old_paragraph': 'NEW_IMAGE_PAGE', 'old_ref': 'NEW', 'new_ref': 'Figure ' + str(n), 'reason': 'Final artwork page heading'})
    output.mkdir(parents=True, exist_ok=True)
    final = output / plan['output_name']
    if final.exists():
        raise FileExistsError('Private output already exists; choose a fresh directory')
    doc.save(final)
    check = Document(final)
    if len(check.inline_shapes) != 6 or len(check.sections) != plan['expected_sections']:
        raise ValueError('Unexpected final image or section count')
    if check.paragraphs[-1]._p.xpath('.//wp:anchor'):
        raise ValueError('Unexpected floating figure anchor')
    with zipfile.ZipFile(final) as package:
        media = [hashlib.sha256(package.read(n)).hexdigest() for n in package.namelist() if n.startswith('word/media/')]
    if any(old_hashes[i] in media for i in plan['reject_old_image_indices']):
        raise ValueError('A rejected original composite remains embedded')
    final_text = '\n'.join(p.text for p in check.paragraphs)
    for phrase in plan['required_final_phrases']:
        if phrase not in final_text:
            raise ValueError('A configured protected phrase disappeared')
    if any(phrase in final_text for phrase in plan['forbidden_final_phrases']):
        raise ValueError('A configured removed phrase persists')
    for spec in plan.get('insertion_order_checks', []):
        if final_text.index(spec['before_text']) >= final_text.index(spec['after_text']):
            raise ValueError('A configured insertion-order assertion failed')
    for paragraph in check.paragraphs:
        for reference in extract(paragraph.text):
            for target in reference['items']:
                if str(target.figure) not in plan['max_panel'] or target.panel > plan['max_panel'][str(target.figure)]:
                    raise ValueError('Reference to missing final main panel')
    changed = {c['old_paragraph'] for c in changes if isinstance(c['old_paragraph'], int)}
    unchanged = []
    for i in range(plan['old_images_start']):
        if i in changed:
            continue
        position = i + sum(item['before_index'] <= i for item in inserted)
        if check.paragraphs[position]._p.xml != before_xml[i]:
            raise ValueError('Unexpected OOXML change in unedited paragraph ' + str(i))
        unchanged.append(i)
    # Exact preserved-reference XML and order, independent of inserted prose.
    for start, stop in plan['preserve_ranges']:
        positions = [i + sum(item['before_index'] <= i for item in inserted) for i in range(start, stop)]
        if [check.paragraphs[i]._p.xml for i in positions] != [before_xml[i] for i in range(start, stop)]:
            raise ValueError('A preserved paragraph range changed')
    if len(references) != sum(len(extract(p.text)) for p in check.paragraphs):
        raise ValueError('Reference ledger does not cover every final occurrence')
    if sha(source) != plan['source_sha256']:
        raise ValueError('Read-only source changed')
    tsv(output / 'FIGURE_REFERENCE_AUDIT.tsv', references, ['old_paragraph', 'old_ref', 'new_ref', 'reason', 'old_targets', 'new_targets'])
    tsv(output / 'PARAGRAPH_CHANGES.tsv', changes, ['old_paragraph', 'operation', 'old_text', 'new_text', 'numeric_guard'])
    tsv(output / 'DELETED_PANEL_LEGEND_SPANS.tsv', deleted, ['old_paragraph', 'old_panel', 'removed_text', 'reason'])
    tsv(output / 'CONFIGURED_TEXT_REPLACEMENTS.tsv', replacements, ['old_paragraph', 'old', 'new', 'reason'])
    (output / 'NUMERIC_GUARD_AUDIT.json').write_text(json.dumps(guards, indent=2), encoding='utf-8')
    section_spec = plan['legend_section']
    inside, legends = False, []
    for paragraph in check.paragraphs:
        if paragraph.text == section_spec['heading']:
            inside = True
            continue
        if inside and paragraph.text == section_spec['end_heading']:
            break
        if inside and paragraph.text:
            legends.append(paragraph.text)
    (output / 'ALL_MAIN_FIGURE_LEGENDS.txt').write_text('\n\n'.join(legends) + '\n', encoding='utf-8')
    receipt = {'source_sha256': sha(source), 'output_sha256': sha(final), 'main_figures': 6,
               'reference_mappings': len(references), 'source_input_unchanged': True, 'numeric_guard': 'PASS',
               'supplementary_references': 'UNCHANGED', 'untouched_paragraphs_xml_verified': unchanged,
               'old_rejected_composites_absent': True, 'embedded_figures': embedded,
               'visual_QA': 'PENDING_ACTUAL_WORD_PDF_AND_EVERY_PAGE_REVIEW'}
    (output / 'MANUSCRIPT_EDIT_RECEIPT.json').write_text(json.dumps(receipt, indent=2), encoding='utf-8')
    print(json.dumps({'output_sha256': receipt['output_sha256'], 'main_figures': 6,
                      'reference_mappings': len(references), 'visual_QA': receipt['visual_QA']}))


if __name__ == '__main__':
    main()
