"""Synthetic technical tests; no research inputs or manuscript text are used."""
import csv
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from output_guard import require_private_output

HERE = Path(__file__).resolve().parent
OPTIONAL_READY = all(importlib.util.find_spec(name) is not None for name in ('docx', 'lxml', 'PIL'))


class PrivateOutputGuard(unittest.TestCase):
    def test_publication_checkout_rejected(self):
        root = next((p for p in HERE.parents if (p / '.git').exists()), None)
        if root is None:
            self.skipTest('Source copy has no Git checkout boundary')
        with self.assertRaisesRegex(ValueError, 'private output'):
            require_private_output(root / 'must-not-write.json')

    def test_external_temp_destination_accepted(self):
        with tempfile.TemporaryDirectory() as directory:
            self.assertEqual(require_private_output(Path(directory) / 'private.json'),
                             (Path(directory) / 'private.json').resolve())


@unittest.skipUnless(OPTIONAL_READY, 'Optional DOCX integration needs python-docx, lxml and Pillow; core mapping tests still run')
class SyntheticPrivateEditing(unittest.TestCase):
    def setUp(self):
        from docx import Document
        from docx.enum.section import WD_SECTION_START
        from PIL import Image
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.inputs = self.root / 'input'
        self.inputs.mkdir()
        self.output = self.root / 'private-output'
        self.source = self.inputs / 'synthetic.docx'
        self.plan_path = self.inputs / 'plan.json'
        self.manifest = self.inputs / 'figures.json'
        self.legend = self.inputs / 'legend.txt'
        source = Document()
        values = ['Synthetic document', 'Methods',
                  'Numeric invariant: 12 units (Figure 1B–E); Supplementary Figure S2.',
                  'Results', 'Synthetic reference to Figure 5B.', 'Figure legends',
                  'Figure 1 | Synthetic first. (A) Remove this box. (B) Item one. (C) Item two. (D) Item three. (E) Item four.',
                  'Figure 2 | Synthetic second.', 'Figure 3 | Synthetic third.',
                  'Figure 4 | Synthetic fourth.',
                  'Figure 5 | Synthetic fifth. (A) Remove this scheme. (B) Item one. (C) Item two. (D) Item three. (E) Item four. (F) Item five. (G) Item six.',
                  'References', 'Synthetic reference entry.']
        for value in values:
            source.add_paragraph(value)
        source.add_section(WD_SECTION_START.NEW_PAGE)
        self.original_images = []
        for i in range(1, 6):
            image = self.inputs / ('old-' + str(i) + '.png')
            Image.new('RGB', (20 + i, 30), (i * 30, 20, 50)).save(image)
            self.original_images.append(hashlib.sha256(image.read_bytes()).hexdigest())
            source.add_paragraph('Figure ' + str(i))
            source.add_paragraph().add_run().add_picture(str(image))
        source.save(self.source)
        self.source_hash = hashlib.sha256(self.source.read_bytes()).hexdigest()
        figures = {}
        for i in range(1, 7):
            image = self.inputs / ('new-' + str(i) + '.png')
            Image.new('RGB', (40 + i, 50), (20, i * 25, 80)).save(image)
            figures[str(i)] = str(image)
        self.manifest.write_text(json.dumps(figures), encoding='utf-8')
        self.legend.write_text('Figure 1 | Synthetic overview. (A) Alpha. (B) Beta. (C) Gamma. (D) Delta. (E) Epsilon.', encoding='utf-8')
        self.plan = {
            'source_sha256': self.source_hash, 'expected_paragraph_count': 24,
            'edit_stop': 13, 'preserve_ranges': [[11, 13]],
            'structure_checks': [{'index': 2, 'starts_with': 'Numeric invariant:'}],
            'panel_edits': [
                {'index': 6, 'remove_panel': 'A', 'old_panel': '1A', 'reason': 'Synthetic layout update',
                 'panel_relabel': {'B': 'A', 'C': 'B', 'D': 'C', 'E': 'D'}},
                {'index': 10, 'remove_panel': 'A', 'old_panel': '5A', 'reason': 'Synthetic layout update',
                 'panel_relabel': {'B': 'A', 'C': 'B', 'D': 'C', 'E': 'D', 'F': 'E', 'G': 'F'}}],
            'reference_overrides': {}, 'text_replacements': {},
            'insertions': [{'before_index': 4, 'source': 'text', 'text': 'Synthetic overview (Figure 1).'},
                           {'before_index': 6, 'source': 'figure1_legend'}],
            'old_images_start': 14, 'expected_old_images': 5,
            'reject_old_image_indices': [0, -1], 'expected_sections': 2,
            'required_final_phrases': ['Numeric invariant: 12 units', 'Supplementary Figure S2.'],
            'forbidden_final_phrases': ['Remove this box.', 'Remove this scheme.'],
            'max_panel': {'1': 'E', '2': 'D', '3': 'F', '4': 'F', '5': 'F', '6': 'F'},
            'legend_section': {'heading': 'Figure legends', 'end_heading': 'References'},
            'insertion_order_checks': [{'before_text': 'Synthetic overview', 'after_text': 'Synthetic reference to Figure'}],
        }

    def tearDown(self):
        self.temp.cleanup()

    def run_editor(self):
        self.plan_path.write_text(json.dumps(self.plan), encoding='utf-8')
        return subprocess.run([sys.executable, '-B', str(HERE / 'revise_manuscript.py'),
                               '--source', str(self.source), '--output-dir', str(self.output),
                               '--plan', str(self.plan_path), '--figure-manifest', str(self.manifest),
                               '--figure1-legend', str(self.legend)], capture_output=True, text=True)

    def test_end_to_end_edit_inventory_and_layout(self):
        from docx import Document
        result = self.run_editor()
        self.assertEqual(result.returncode, 0, result.stderr)
        final = self.output / 'manuscript_renumbered.docx'
        document = Document(final)
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('12 units (Figure 2A–D)', text)
        self.assertIn('Synthetic reference to Figure 6A.', text)
        self.assertIn('Supplementary Figure S2.', text)
        self.assertNotIn('Remove this box.', text)
        self.assertEqual(len(document.inline_shapes), 6)
        self.assertEqual(len(document.sections), 2)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.source_hash)
        receipt = json.loads((self.output / 'MANUSCRIPT_EDIT_RECEIPT.json').read_text())
        self.assertEqual(receipt['numeric_guard'], 'PASS')
        self.assertTrue(receipt['old_rejected_composites_absent'])
        self.assertIn('PENDING', receipt['visual_QA'])
        with (self.output / 'FIGURE_REFERENCE_AUDIT.tsv').open(encoding='utf-8-sig') as stream:
            self.assertEqual(len(list(csv.DictReader(stream, delimiter='\t'))), receipt['reference_mappings'])
        inventory = self.root / 'private-inventory.json'
        run = subprocess.run([sys.executable, '-B', str(HERE / 'inspect_docx.py'), str(final), '--output', str(inventory)], capture_output=True, text=True)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertTrue(json.loads(inventory.read_text())['source_input_unchanged'])
        layout_output = self.root / 'private-layout.docx'
        layout_audit = self.root / 'private-layout.json'
        before = final.read_bytes()
        run = subprocess.run([sys.executable, '-B', str(HERE / 'fix_legend_pagination.py'),
                              '--source', str(final), '--output', str(layout_output),
                              '--expected-sha256', hashlib.sha256(before).hexdigest(),
                              '--audit', str(layout_audit)], capture_output=True, text=True)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertEqual(final.read_bytes(), before)
        self.assertEqual([p.text for p in Document(layout_output).paragraphs], [p.text for p in document.paragraphs])
        audit = json.loads(layout_audit.read_text())
        self.assertTrue(audit['source_input_unchanged'])
        self.assertEqual(audit['all_media_and_other_package_parts'], 'EXACTLY_UNCHANGED')

    def test_changed_numeric_token_is_rejected(self):
        self.plan['text_replacements'] = {'2': [{'old': '12 units', 'new': '13 units', 'reason': 'Deliberate synthetic drift'}]}
        result = self.run_editor()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Scientific numeric drift', result.stderr)
        self.assertFalse(self.output.exists())

    def test_source_hash_mismatch_is_rejected(self):
        self.plan['source_sha256'] = '0' * 64
        result = self.run_editor()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Selected source SHA mismatch', result.stderr)
        self.assertFalse(self.output.exists())

    def test_missing_protected_phrase_is_rejected(self):
        self.plan['required_final_phrases'].append('Absent synthetic sentinel')
        result = self.run_editor()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('protected phrase disappeared', result.stderr)
        self.assertFalse((self.output / 'MANUSCRIPT_EDIT_RECEIPT.json').exists())


if __name__ == '__main__':
    unittest.main()
